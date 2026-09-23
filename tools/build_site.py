"""Build the static website (docs/) from the review's files. Usage: python tools/build_site.py
Needs: markdown, openpyxl. Output is plain HTML/CSS/JS with no external dependencies (works on GitHub Pages)."""
import csv
import html
import re
import shutil
from pathlib import Path

import markdown
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
NORTHBRIDGE_REPO = "https://github.com/VAishwaryaSingh/Northbridge-Ledger-Agent"
NORTHBRIDGE_DEMO = "https://northbridge-ledger-agent.streamlit.app/"
ASSURANCE_REPO = "https://github.com/VAishwaryaSingh/northbridge-ai-assurance"
COMMIT = "9c4127ad65aa5a9e7d362c2d05f69ecfdbf7c5f9"

NAV = [("index.html", "Home"), ("findings.html", "Findings"), ("method.html", "Method"), ("testing.html", "Testing"),
       ("remediation.html", "Remediation"), ("toolkit.html", "Toolkit"), ("downloads/assurance_report.pdf", "Report (PDF)")]

CSS = """
:root{--bg:#fff;--fg:#1b1f27;--muted:#5a6270;--line:#d9dee6;--soft:#f3f5f9;--brand:#1f3864;--link:#1a5fb4;--bad:#b71c1c;--badbg:#fdecea;--warn:#8a5a00;--warnbg:#fff4d6;--ok:#1b6e34;--okbg:#e6f4ea;--code:#eef1f6}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--bg:#0f1420;--fg:#e6e9ef;--muted:#a3abb8;--line:#2b3446;--soft:#171e2e;--brand:#9db8ee;--link:#8ab4f8;--bad:#ff8a80;--badbg:#3a1a1a;--warn:#ffd27a;--warnbg:#3a2e12;--ok:#7fd99a;--okbg:#14301f;--code:#1d2536}}
:root[data-theme=dark]{--bg:#0f1420;--fg:#e6e9ef;--muted:#a3abb8;--line:#2b3446;--soft:#171e2e;--brand:#9db8ee;--link:#8ab4f8;--bad:#ff8a80;--badbg:#3a1a1a;--warn:#ffd27a;--warnbg:#3a2e12;--ok:#7fd99a;--okbg:#14301f;--code:#1d2536}
*{box-sizing:border-box}html{-webkit-text-size-adjust:100%}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.6 system-ui,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif}
a{color:var(--link)}a:focus-visible,button:focus-visible,input:focus-visible{outline:3px solid var(--link);outline-offset:2px}
.skip{position:absolute;left:-999px;top:0;background:var(--brand);color:#fff;padding:8px 12px;z-index:10}.skip:focus{left:8px;top:8px}
header.site{border-bottom:1px solid var(--line);background:var(--bg);position:sticky;top:0;z-index:5}
.wrap{max-width:920px;margin:0 auto;padding:0 16px}
header.site .wrap{display:flex;flex-wrap:wrap;align-items:center;gap:4px 14px;min-height:52px}
.brand{font-weight:700;color:var(--brand);text-decoration:none;margin-right:6px}
nav{display:flex;flex-wrap:wrap;gap:2px 4px}nav a{padding:6px 9px;border-radius:6px;text-decoration:none;color:var(--fg);font-size:.92rem}
nav a:hover{background:var(--soft)}nav a[aria-current=page]{background:var(--soft);font-weight:600;box-shadow:inset 0 -2px 0 var(--brand)}
.banner{background:var(--warnbg);color:var(--fg);border-bottom:1px solid var(--line);font-size:.88rem}
.banner .wrap{padding-top:8px;padding-bottom:8px}.banner strong{color:var(--warn)}
main{padding:24px 0 56px}h1{font-size:1.9rem;line-height:1.25;color:var(--brand);margin:.2em 0 .5em}h2{font-size:1.35rem;color:var(--brand);margin:1.8em 0 .5em;padding-top:.3em}h3{font-size:1.1rem;margin:1.4em 0 .4em}
p,li{max-width:74ch}code{background:var(--code);padding:1px 5px;border-radius:4px;font-size:.88em}pre{background:var(--code);padding:12px;border-radius:8px;overflow:auto}pre code{background:none;padding:0}
.scroll{overflow-x:auto;margin:12px 0;-webkit-overflow-scrolling:touch}table{border-collapse:collapse;width:100%;font-size:.9rem}
th,td{border:1px solid var(--line);padding:6px 9px;text-align:left;vertical-align:top}th{background:var(--brand);color:#fff}@media (prefers-color-scheme:dark){:root:not([data-theme=light]) th{color:#0f1420}}
tbody tr:nth-child(even){background:var(--soft)}
.chip{display:inline-block;padding:1px 9px;border-radius:99px;font-size:.78rem;font-weight:700;white-space:nowrap}
.High,.notmet{background:var(--badbg);color:var(--bad)}.Medium,.warn{background:var(--warnbg);color:var(--warn)}.Low,.met{background:var(--okbg);color:var(--ok)}
.hero{background:var(--soft);border:1px solid var(--line);border-radius:12px;padding:20px 22px;margin:8px 0 22px}.hero p{max-width:none}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:12px;margin:14px 0}
.card{border:1px solid var(--line);border-radius:10px;padding:12px 14px;background:var(--bg)}.card h3{margin:.1em 0 .3em;font-size:1rem}.card p{margin:.2em 0;font-size:.92rem;color:var(--muted)}
.stat{font-size:1.6rem;font-weight:700;color:var(--brand);line-height:1.2}
.filter{width:100%;max-width:420px;padding:8px 10px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg);font:inherit;margin:6px 0}
.chart text{fill:var(--fg);font:12px system-ui,sans-serif}.chart .muted{fill:var(--muted)}
footer{border-top:1px solid var(--line);color:var(--muted);font-size:.85rem;padding:18px 0 34px}
.toc{background:var(--soft);border:1px solid var(--line);border-radius:8px;padding:6px 16px;margin:14px 0}.toc ul{margin:.3em 0;padding-left:1.2em}
button.theme{margin-left:auto;background:none;border:1px solid var(--line);border-radius:6px;color:var(--fg);padding:4px 9px;cursor:pointer;font:inherit;font-size:.85rem}
@media (max-width:560px){h1{font-size:1.5rem}body{font-size:15.5px}}
@media print{header.site,.filter,button.theme{display:none}body{font-size:11pt}}
"""

JS = """
(function(){var r=document.documentElement,t=null;try{t=localStorage.getItem('theme')}catch(e){}if(t)r.setAttribute('data-theme',t);
var b=document.getElementById('themebtn');if(b)b.addEventListener('click',function(){var d=r.getAttribute('data-theme')||(matchMedia('(prefers-color-scheme:dark)').matches?'dark':'light');var n=d==='dark'?'light':'dark';r.setAttribute('data-theme',n);try{localStorage.setItem('theme',n)}catch(e){}});
document.querySelectorAll('input.filter').forEach(function(i){i.addEventListener('input',function(){var q=i.value.toLowerCase(),tb=document.getElementById(i.dataset.target);tb.querySelectorAll('tbody tr').forEach(function(tr){tr.style.display=tr.textContent.toLowerCase().indexOf(q)>-1?'':'none'})})});})();
"""


def slug(value, sep="-"):
    m = re.match(r"\s*(F-\d\d)\b", value)
    if m:
        return m.group(1).lower()
    return re.sub(r"[^a-z0-9]+", sep, value.lower()).strip(sep)


def md_to_html(text, toc=False):
    md = markdown.Markdown(extensions=["tables", "fenced_code", "sane_lists", "toc"], extension_configs={"toc": {"slugify": slug}})
    out = md.convert(text)
    out = re.sub(r"<table>", '<div class="scroll"><table>', out)
    out = re.sub(r"</table>", "</table></div>", out)
    return out, md


def chip(word):
    cls = {"High": "High", "Medium": "Medium", "Low": "Low"}.get(word, "")
    return f'<span class="chip {cls}">{word}</span>'


def page(fname, title, body, description=""):
    depth = fname.count("/")
    pre = "../" * depth
    nav = "".join(f'<a href="{pre}{h}"' + (' aria-current="page"' if h == fname else "") + f">{html.escape(l)}</a>" for h, l in NAV)
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)} - Northbridge Assurance Review</title>
<meta name="description" content="{html.escape(description or 'Self-assessment of a ledger anomaly-detection tool using audit methodology.')}">
<style>{CSS}</style></head><body>
<a class="skip" href="#main">Skip to content</a>
<header class="site"><div class="wrap"><a class="brand" href="{pre}index.html">Northbridge Assurance Review</a><nav aria-label="Main">{nav}</nav><button class="theme" id="themebtn" type="button">Light / dark</button></div></header>
<div class="banner"><div class="wrap"><strong>Self-assessment, not an independent audit.</strong> The reviewer also built the system. Figures describe a constructed test set and are not estimates of accuracy on real books. No human independent review has been done.</div></div>
<main id="main"><div class="wrap">
{body}
</div></main>
<footer><div class="wrap"><a href="{ASSURANCE_REPO}">Source, scripts and full evidence on GitHub</a>. Subject: <a href="{NORTHBRIDGE_REPO}">Northbridge Ledger Agent</a>, commit <code>{COMMIT[:7]}</code> (23 Sep 2026). Live demo: <a href="{NORTHBRIDGE_DEMO}">northbridge-ledger-agent.streamlit.app</a>. Synthetic and sandbox data only. The Northbridge project itself was not changed. This page states a conclusion, not an audit opinion, and gives no certification.</div></footer>
<script>{JS}</script></body></html>"""


def write(fname, title, body, description=""):
    p = DOCS / fname
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(page(fname, title, body, description), encoding="utf-8")


def read(rel):
    return (ROOT / rel).read_text(encoding="utf-8")


def read_csv(rel):
    return list(csv.DictReader(open(ROOT / rel, encoding="utf-8")))


# ------------------------------------------------------------------ setup
if DOCS.exists():
    shutil.rmtree(DOCS)
(DOCS / "downloads").mkdir(parents=True)
(DOCS / ".nojekyll").write_text("")
DL = [("05_reporting/assurance_report.pdf", "Assurance report (PDF, 12 pages)"), ("05_reporting/findings.md", "Findings (Markdown)"),
      ("05_reporting/remediation.patch", "Remediation patch (not applied to Northbridge)"), ("02_risk/risk_register.xlsx", "Risk register (Excel)"),
      ("03_controls/control_matrix.xlsx", "Control matrix (Excel)"), ("toolkit/ai_audit_checklist.md", "AI audit checklist (Markdown)"),
      ("toolkit/ai_readiness_review.md", "AI readiness template (Markdown)"), ("04_testing/test_dataset/labels.csv", "Test set ground-truth labels (CSV)"),
      ("04_testing/test_dataset/ledger_test.db", "Synthetic test ledger (SQLite)"), ("04_testing/results/criteria_assessment_full.csv", "Criteria assessment (CSV)"),
      ("04_testing/results/metrics_by_error_type.csv", "Metrics by error type (CSV)"), ("04_testing/results/sensitivity.csv", "Labelling sensitivity (CSV)"),
      ("04_testing/results/item_results.csv", "Item-level results (CSV)"), ("04_testing/evidence_log.md", "Evidence log (Markdown)")]
for rel, _ in DL:
    shutil.copy(ROOT / rel, DOCS / "downloads" / Path(rel).name)

# ------------------------------------------------------------------ data for home
crit = read_csv("04_testing/results/criteria_assessment_full.csv")
before = {r["error_type"]: r for r in read_csv("04_testing/results/metrics_by_error_type.csv")}
after = {r["error_type"]: r for r in read_csv("04_testing/results_remediated/metrics_by_error_type.csv")}

def status_chip(a):
    if a.startswith("NOT MET"):
        return '<span class="chip notmet">Not met</span>' + (" borderline" if "borderline" in a else "")
    return '<span class="chip met">Met on point estimate</span>'

crit_rows = "".join(f"<tr><td>{html.escape(c['criterion'])}</td><td>{html.escape(c['threshold'])}</td><td>{html.escape(c['result'])}</td><td>{status_chip(c['assessment'])}</td></tr>" for c in crit)

names = {"DUP_AR": "Duplicate sales invoice", "DUP_AP": "Duplicate supplier bill", "DUP_PAY": "Duplicate payment", "BREAK": "Reconciliation break",
         "VAT_MISCODE": "VAT miscoding", "THRESHOLD": "Just under threshold", "WEEKEND": "Weekend-dated bill", "OUTLIER": "Outlier amount"}
types = list(names)
W, LM, BW = 700, 190, 420
H = 40 + len(types) * 34
bars = []
for i, t in enumerate(types):
    y = 30 + i * 34
    b, a = float(before[t]["recall_any_rule"]), float(after[t]["recall_any_rule"])
    thr = 0.95 if t in ("DUP_AR", "DUP_AP", "DUP_PAY", "BREAK") else 0.85
    bars.append(f'<text x="{LM-8}" y="{y+14}" text-anchor="end">{names[t]} (n={before[t]["n_errors"]})</text>'
                f'<rect x="{LM}" y="{y}" width="{max(b*BW,1):.1f}" height="11" fill="#9aa3af"/>'
                f'<rect x="{LM}" y="{y+13}" width="{max(a*BW,1):.1f}" height="11" fill="#1f5fbf"/>'
                f'<text x="{LM+max(a,b)*BW+6:.1f}" y="{y+16}" class="muted">{b:.0%} / {a:.0%}</text>'
                f'<line x1="{LM+thr*BW:.1f}" x2="{LM+thr*BW:.1f}" y1="{y-3}" y2="{y+27}" stroke="currentColor" stroke-dasharray="3,3"/>')
chart = (f'<svg class="chart" role="img" aria-label="Recall by error type before and after remediation. Duplicate supplier bills rose from 0 to 100 percent; duplicate payment fell from 100 to 75 percent." '
         f'viewBox="0 0 {W} {H}" width="100%" style="max-width:{W}px"><text x="0" y="14" style="font-weight:700">Recall by error type: before (grey) and after remediation (blue). Dashed line = criterion.</text>{"".join(bars)}</svg>')

F = [("F-01", "High", "Duplicate supplier bills are not detected"), ("F-02", "High", "Reconciliation output depends on record order and mis-attributes recurring payments"),
     ("F-03", "High", "No run log or audit trail"), ("F-04", "High", "No human disposition of flags; no owner or intended-use statement"),
     ("F-05", "High", "Reported accuracy rests on circular validation and a static score"), ("F-06", "Medium", "No regression tests, version tags or pinned dependencies"),
     ("F-07", "Medium", "Variance explanations can point against the headline movement"), ("F-08", "Medium", "False positives on legitimate items exceed the criterion"),
     ("F-09", "Medium", "Excess QuickBooks permission and plaintext refresh tokens"), ("F-10", "Medium", 'Ambiguous "AI-native agent" positioning and missing limitations statement'),
     ("F-11", "Medium", "Data completeness is not evidenced"), ("F-12", "Low", "Hard-coded, unapproved thresholds and message defects"),
     ("F-13", "Low", "QuickBooks bank purchases all reported as unmatched")]
flist = "".join(f'<tr><td><a href="findings.html#{i.lower()}">{i}</a></td><td>{html.escape(t)}</td><td>{chip(r)}</td></tr>' for i, r, t in F)

home = f"""
<div class="hero"><h1>Is a "7 out of 7" result evidence that an AI finance tool can be trusted?</h1>
<p>This is a structured assurance review of the <strong>Northbridge Ledger Agent</strong>, a tool that flags reconciliation breaks, duplicates and anomalies in Xero and QuickBooks data. Its own validation planted seven errors and caught all seven. This review applies audit methodology to find out what that result does and does not show.</p>
<p><strong>Conclusion.</strong> As at the pinned commit, Northbridge should <strong>not be relied on as a control</strong> for month-end reconciliation. It works as a demonstration and a lead generator for the error types it targets. Of eight criteria set before testing, seven were not met. The tool is titled an "Agent" and positioned as a small-scale version of AI-native products, but calls <strong>no LLM</strong>: it is deterministic rules plus a statistical check.</p>
<p><a href="downloads/assurance_report.pdf"><strong>Read the 12-page report (PDF)</strong></a> &middot; <a href="findings.html">See the findings</a> &middot; <a href="toolkit.html">Use the toolkit</a></p></div>
<div class="grid">
<div class="card"><div class="stat">230</div><h3>test items</h3><p>31 real errors, 199 clean items, 32 of them legitimate look-alikes. Synthetic data only.</p></div>
<div class="card"><div class="stat">13</div><h3>findings</h3><p>5 High, 6 Medium, 2 Low, each in Condition, Criteria, Cause, Effect, Recommendation form.</p></div>
<div class="card"><div class="stat">7 of 8</div><h3>criteria not met</h3><p>Fixed before testing. Recall on duplicates and breaks 83%, precision 56%.</p></div>
<div class="card"><div class="stat">0 of 3</div><h3>duplicate supplier bills caught</h3><p>The duplicate check only looks at sales invoices. A gap the 7/7 result never showed.</p></div>
</div>
<h2 id="results">Results against the criteria</h2>
<p>Measured on a test set built by the reviewer after reading the rules, so it is <strong>not independent of the rules</strong>. It shows how the system behaves on that mix, not how it will behave on real books.</p>
<div class="scroll"><table><thead><tr><th>Criterion</th><th>Threshold</th><th>Result</th><th>Assessment</th></tr></thead><tbody>{crit_rows}</tbody></table></div>
<h2 id="chart">Recall by error type, before and after remediation</h2>
{chart}
<p>Fixes were tried on a copy and delivered as a patch; the Northbridge project itself was not changed. Two costs are on the <a href="remediation.html">remediation page</a>: one duplicate payment is no longer flagged as such, and the project's published synthetic precision falls from 0.778 to 0.712. The fixes were designed against the same dataset used to retest, so they do not show real-world accuracy has improved.</p>
<h2 id="findings">The 13 findings</h2>
<div class="scroll"><table><thead><tr><th>ID</th><th>Finding</th><th>Rating</th></tr></thead><tbody>{flist}</tbody></table></div>
<h2 id="how">How the review was done</h2>
<ol><li><strong>Planning:</strong> terms of reference with criteria fixed before testing, and a system description showing where data flows.</li>
<li><strong>Risk:</strong> 28 risks scored for likelihood and impact and mapped to frameworks. LLM-only risks marked not applicable.</li>
<li><strong>Controls:</strong> 29 controls assessed for design; missing controls recorded as gaps and not added before testing.</li>
<li><strong>Testing:</strong> a synthetic ledger with labels frozen by hash before the first run, an evaluation harness against a read-only copy of the pinned commit, and control tests.</li>
<li><strong>Challenge:</strong> a separate Claude session reviewed the raw evidence and raised 12 challenges, all accepted. It found the tool's output changes with record order, which the reviewer had missed. This is a model-based challenge, not a human review.</li>
<li><strong>Findings, remediation and report:</strong> fixes demonstrated on a copy and retested, with costs stated.</li></ol>
<h2 id="limits">What this does not show</h2>
<ul><li>Accuracy on real books: the data is synthetic and the test set was designed after reading the rules.</li>
<li>Small samples: 3 to 6 errors per type, and false positives cluster in about six causes.</li>
<li>Not tested: live connectors and pagination, portal-granted permissions, multi-line bills, voided or null data, multi-currency.</li>
<li>Framework references (NIST, ISO) are at area level and need checking against the source documents; EU AI Act dates were not verified.</li>
<li>No independent human review has been done.</li></ul>
<h2 id="reuse">Reuse it</h2>
<p>The <a href="toolkit.html">toolkit</a> contains a checklist for assessing an AI system and a template for assessing a process for AI adoption. Both were built from this one review and are not validated on other systems.</p>
"""
write("index.html", "Home", home, "A self-assessment of an AI-described finance tool: seven of eight criteria not met, 13 findings, remediation on a copy, and a reusable toolkit.")

# ------------------------------------------------------------------ findings
ftxt = read("05_reporting/findings.md")
body, _ = md_to_html(ftxt)
body = re.sub(r"<h2 id=\"(f-\d\d)\">(F-\d\d [^<]*?) \((High|Medium|Low)\)</h2>", lambda m: f'<h2 id="{m.group(1)}">{m.group(2)} {chip(m.group(3))}</h2>', body)
body = re.sub(r"<td>(High|Medium|Low)</td>", lambda m: f"<td>{chip(m.group(1))}</td>", body)
write("findings.html", "Findings", body + '<p><a href="downloads/findings.md">Download findings (Markdown)</a></p>')

# ------------------------------------------------------------------ method pages
def md_page(fname, title, rel, extra=""):
    text = read(rel)
    text = re.sub(r"```mermaid.*?```", "@@DIAGRAM@@", text, flags=re.S)
    b, _ = md_to_html(text)
    b = b.replace("<p>@@DIAGRAM@@</p>", DIAGRAM)
    write(fname, title, b + extra)

DIAGRAM = """<figure><svg class="chart" role="img" aria-label="Data flow: Xero and QuickBooks APIs to connectors, local SQLite, matcher and rules, dashboard; demo databases published to GitHub and Streamlit Cloud" viewBox="0 0 720 250" width="100%" style="max-width:720px">
<g font-size="12" font-family="system-ui,sans-serif">
<rect x="0" y="20" width="120" height="56" rx="6" fill="#fdecea" stroke="#b71c1c"/><text x="60" y="44" text-anchor="middle">Xero API</text><text x="60" y="60" text-anchor="middle">read scopes</text>
<rect x="0" y="110" width="120" height="70" rx="6" fill="#fdecea" stroke="#b71c1c"/><text x="60" y="134" text-anchor="middle">QuickBooks API</text><text x="60" y="150" text-anchor="middle">read/write scope</text><text x="60" y="166" text-anchor="middle">granted</text>
<rect x="170" y="50" width="130" height="90" rx="6" fill="#e8eef7" stroke="#1f3864"/><text x="235" y="76" text-anchor="middle" fill="#1b1f27">Connectors</text><text x="235" y="94" text-anchor="middle" fill="#1b1f27">OAuth + GET</text><text x="235" y="110" text-anchor="middle" fill="#1b1f27">.env: keys,</text><text x="235" y="126" text-anchor="middle" fill="#1b1f27">plaintext tokens</text>
<rect x="340" y="50" width="100" height="90" rx="6" fill="#e8eef7" stroke="#1f3864"/><text x="390" y="90" text-anchor="middle" fill="#1b1f27">SQLite</text><text x="390" y="108" text-anchor="middle" fill="#1b1f27">ledger (local)</text>
<rect x="480" y="50" width="110" height="90" rx="6" fill="#e8eef7" stroke="#1f3864"/><text x="535" y="84" text-anchor="middle" fill="#1b1f27">Matcher, rules</text><text x="535" y="102" text-anchor="middle" fill="#1b1f27">A1-A7, variance</text><text x="535" y="120" text-anchor="middle" fill="#1b1f27">templates</text>
<rect x="620" y="50" width="100" height="90" rx="6" fill="#e8eef7" stroke="#1f3864"/><text x="670" y="86" text-anchor="middle" fill="#1b1f27">Dashboard</text><text x="670" y="104" text-anchor="middle" fill="#1b1f27">display only,</text><text x="670" y="120" text-anchor="middle" fill="#1b1f27">no sign-off</text>
<rect x="340" y="190" width="200" height="46" rx="6" fill="#fdecea" stroke="#b71c1c"/><text x="440" y="210" text-anchor="middle" fill="#1b1f27">Public GitHub +</text><text x="440" y="226" text-anchor="middle" fill="#1b1f27">Streamlit Cloud (demo copies)</text>
<g stroke="currentColor" stroke-width="1.5" fill="none"><path d="M120 48 L170 80"/><path d="M120 145 L170 115"/><path d="M300 95 L340 95"/><path d="M440 95 L480 95"/><path d="M590 95 L620 95"/><path d="M390 140 L410 190"/></g></g></svg>
<figcaption>Red boxes are where data leaves the developer machine. No LLM provider receives data.</figcaption></figure>"""

md_page("terms.html", "Terms of reference", "01_planning/terms_of_reference.md")
md_page("system.html", "System description", "01_planning/system_description.md")
md_page("regulatory.html", "Regulatory classification", "02_risk/ai_act_classification.md")
md_page("fs-audit.html", "Financial statement audit implications", "05_reporting/fs_audit_implications.md")

# risks
rr = load_workbook(ROOT / "02_risk/risk_register.xlsx")["Risk Register"]
rows = []
for r in range(2, rr.max_row + 1):
    l, i = rr.cell(r, 4).value, rr.cell(r, 5).value
    s = l * i
    rating = "High" if s >= 15 else "Medium" if s >= 8 else "Low"
    rows.append((rr.cell(r, 1).value, rr.cell(r, 2).value, rr.cell(r, 3).value, l, i, s, rating, rr.cell(r, 8).value, rr.cell(r, 9).value, rr.cell(r, 10).value, rr.cell(r, 12).value, rr.cell(r, 13).value))
rt = "".join("<tr>" + "".join(f"<td>{html.escape(str(x))}</td>" if k != 6 else f"<td>{chip(x)}</td>" for k, x in enumerate(row)) + "</tr>" for row in rows)
risks = f"""<h1>Risk register</h1><p>28 risks scored 1 to 5 for likelihood and impact <strong>before</strong> controls. Inherent score is the product: High 15 or more, Medium 8 to 14, Low below 8. Scores are the reviewer's judgement. NIST and ISO references are at area level and must be checked against the source documents.</p>
<p><a href="downloads/risk_register.xlsx">Download the workbook (with scoring rationale)</a></p>
<label for="rf">Filter risks</label><br><input class="filter" id="rf" type="search" data-target="rt" placeholder="Type to filter, for example: security, R08, High">
<div class="scroll"><table id="rt"><thead><tr><th>ID</th><th>Risk</th><th>Category</th><th>L</th><th>I</th><th>Score</th><th>Rating</th><th>NIST function</th><th>NIST ref</th><th>ISO/IEC 42001</th><th>UK principle</th><th>Basis</th></tr></thead><tbody>{rt}</tbody></table></div>
<h2>Not applicable (no LLM is called)</h2><ul><li>Prompt injection via transaction descriptions</li><li>Hallucinated or unsupported explanations</li><li>Provider data handling, retention and training terms</li><li>Silent model version changes; model and prompt inventory</li></ul>"""
write("risks.html", "Risk register", risks)

cm = load_workbook(ROOT / "03_controls/control_matrix.xlsx")["Control Matrix"]
crow = []
for r in range(2, cm.max_row + 1):
    v = [cm.cell(r, c).value or "" for c in (1, 2, 3, 6, 7, 11, 9)]
    ip = v[3]
    ipchip = f'<span class="chip {"Low" if ip == "Y" else "Medium" if ip == "Partial" else "High"}">{ip}</span>'
    crow.append(f"<tr><td>{html.escape(v[0])}</td><td>{html.escape(v[1])}</td><td>{html.escape(v[2])}</td><td>{ipchip}</td><td>{html.escape(v[4])}</td><td>{html.escape(v[5])}</td><td>{html.escape(v[6])}</td></tr>")
controls = f"""<h1>Control matrix</h1><p>29 controls identified from the code and documents at the pinned commit. <strong>Y</strong> = design adequate, <strong>Partial</strong> = design limited, <strong>N</strong> = not in place. Missing controls were recorded as gaps and were not added before testing. Operating results come from the Phase 4 tests and were corrected after the challenge review.</p>
<p><a href="downloads/control_matrix.xlsx">Download the workbook</a></p>
<label for="cf">Filter controls</label><br><input class="filter" id="cf" type="search" data-target="ct" placeholder="Type to filter, for example: logging, C26, not in place">
<div class="scroll"><table id="ct"><thead><tr><th>ID</th><th>Risk(s)</th><th>Control</th><th>In place</th><th>Design test</th><th>Operating result</th><th>Evidence</th></tr></thead><tbody>{"".join(crow)}</tbody></table></div>"""
write("controls.html", "Control matrix", controls)

method = """<h1>Method</h1><p>How the review was set up, scoped and assessed. Read these in order or jump to what you need.</p>
<div class="grid">
<div class="card"><h3><a href="terms.html">Terms of reference</a></h3><p>Objective, scope, criteria fixed before testing, independence disclosure, limitations, and Amendment 1 made after the challenge review.</p></div>
<div class="card"><h3><a href="system.html">System description</a></h3><p>How Northbridge works, that it calls no LLM, where data flows, permissions, secrets, logging.</p></div>
<div class="card"><h3><a href="regulatory.html">Regulatory classification</a></h3><p>EU AI Act reasoning and the five UK principles.</p></div>
<div class="card"><h3><a href="risks.html">Risk register</a></h3><p>28 scored risks mapped to NIST AI RMF, ISO/IEC 42001 and UK principles.</p></div>
<div class="card"><h3><a href="controls.html">Control matrix</a></h3><p>29 controls with design and operating results.</p></div>
<div class="card"><h3><a href="fs-audit.html">Financial statement audit implications</a></h3><p>How an external auditor would treat the tool (ISA 315, 330, 500, 265).</p></div>
</div>"""
write("method.html", "Method", method)

# ------------------------------------------------------------------ testing
t_body, _ = md_to_html(read("04_testing/evidence_log.md"))
ds, _ = md_to_html(read("04_testing/test_dataset/README.md"))
testing = f"""<h1>Testing</h1><p>Every test was logged with time, method, sample and result. Later entries correct earlier ones where the challenge review found wording that overreached. See also the <a href="challenge.html">challenge review</a>.</p>
<div class="toc"><strong>On this page</strong><ul><li><a href="#evidence-log">Evidence log</a></li><li><a href="#test-dataset">Test dataset</a></li><li><a href="#data-files">Data files</a></li></ul></div>
<h2 id="evidence-log">Evidence log</h2>{t_body.replace('<h1 id="evidence-log">Evidence Log</h1>','')}
<h2 id="test-dataset">Test dataset</h2>{ds.replace('<h1 id="test-dataset">Test dataset</h1>','')}
<h2 id="data-files">Data files</h2><ul><li><a href="downloads/labels.csv">labels.csv</a> (ground truth, 230 items)</li><li><a href="downloads/ledger_test.db">ledger_test.db</a> (synthetic ledger, SQLite)</li><li><a href="downloads/criteria_assessment_full.csv">criteria_assessment_full.csv</a></li><li><a href="downloads/metrics_by_error_type.csv">metrics_by_error_type.csv</a></li><li><a href="downloads/sensitivity.csv">sensitivity.csv</a></li><li><a href="downloads/item_results.csv">item_results.csv</a></li></ul>"""
write("testing.html", "Testing", testing)
c_body, _ = md_to_html(read("04_testing/challenge_review.md"))
write("challenge.html", "Challenge review", c_body + '<p>The raw report from the challenge reviewer is in the repository as <code>04_testing/challenge_review_raw.md</code>.</p>')
r_body, _ = md_to_html(read("05_reporting/remediation.md"))
write("remediation.html", "Remediation", r_body + '<p><a href="downloads/remediation.patch">Download remediation.patch</a> (applies to commit 9c4127a; the Northbridge project itself is unchanged).</p>')

# ------------------------------------------------------------------ toolkit
toolkit = """<h1>Toolkit</h1><p>Two reusable documents. Both were built from this one review plus general practice and have <strong>not been validated on another system</strong>. LLM-specific checks come from framework guidance and have not been exercised in practice.</p>
<div class="grid">
<div class="card"><h3><a href="checklist.html">AI system assurance checklist</a></h3><p>11 sections, each item with a design test and an operating test. Starts with "is it really AI?" and ends with a 30-minute triage list. Lessons from this review are built in.</p></div>
<div class="card"><h3><a href="readiness.html">AI readiness review template</a></h3><p>For assessing a process before adopting AI or automation: process map, baseline, opportunity scoring, data readiness, ROI with visible assumptions, risks and controls.</p></div>
</div>
<p>Downloads: <a href="downloads/ai_audit_checklist.md">checklist (Markdown)</a> &middot; <a href="downloads/ai_readiness_review.md">template (Markdown)</a></p>
<h2>Using them on someone else's system</h2><p>They work as a self-assessment when you have access to the system, its data and its controls. They cannot be answered from a public website alone: accuracy, determinism, logging, human oversight, permissions and change control all need inside access. A public page can only show what a business says about itself.</p>"""
write("toolkit.html", "Toolkit", toolkit)
for fname, title, rel in [("checklist.html", "AI system assurance checklist", "toolkit/ai_audit_checklist.md"), ("readiness.html", "AI readiness review template", "toolkit/ai_readiness_review.md")]:
    b, _ = md_to_html(read(rel))
    write(fname, title, b)

print("built", sum(1 for _ in DOCS.rglob("*.html")), "pages")
