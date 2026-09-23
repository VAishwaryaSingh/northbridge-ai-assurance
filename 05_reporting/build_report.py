"""Build assurance_report.pdf from the review's evidence files. Usage: python build_report.py
Needs reportlab and openpyxl. Numbers are read from the results and workbooks where possible."""
import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "05_reporting" / "assurance_report.pdf"
NAVY, GREY, RED, GREEN, AMBER = colors.HexColor("#1F3864"), colors.HexColor("#595959"), colors.HexColor("#B71C1C"), colors.HexColor("#2E7D32"), colors.HexColor("#B26A00")
LIGHT = colors.HexColor("#F2F4F8")

ss = getSampleStyleSheet()
body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5, alignment=TA_LEFT)
small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, spaceAfter=2)
cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10, spaceAfter=0)
cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
cellw = ParagraphStyle("cellw", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=NAVY, spaceBefore=6, spaceAfter=7)
h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=NAVY, spaceBefore=8, spaceAfter=4)
title = ParagraphStyle("title", parent=body, fontName="Helvetica-Bold", fontSize=24, leading=29, textColor=NAVY, spaceAfter=6)
sub = ParagraphStyle("sub", parent=body, fontSize=12, leading=16, textColor=GREY, spaceAfter=4)
callout = ParagraphStyle("callout", parent=body, fontSize=10, leading=14, backColor=LIGHT, borderPadding=(7, 7, 7, 7), spaceBefore=6, spaceAfter=10)


def P(t, st=body):
    return Paragraph(t, st)


def tbl(rows, widths, header=True, zebra=True, extra=None):
    data = [[c if not isinstance(c, str) else P(c, cellw if (header and i == 0) else cell) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED6")),
             ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if header:
        style.append(("BACKGROUND", (0, 0), (-1, 0), NAVY))
    if zebra:
        for i in range(1 if header else 0, len(rows)):
            if i % 2 == 0:
                style.append(("BACKGROUND", (0, i), (-1, i), LIGHT))
    t.setStyle(TableStyle(style + (extra or [])))
    return t


def read_csv(p):
    return list(csv.DictReader(open(ROOT / p)))


# ---------- data ----------
overall = {r["metric"]: r for r in read_csv("04_testing/results/metrics_overall.csv")}
before_t = {r["error_type"]: r for r in read_csv("04_testing/results/metrics_by_error_type.csv")}
after_t = {r["error_type"]: r for r in read_csv("04_testing/results_remediated/metrics_by_error_type.csv")}
sens = read_csv("04_testing/results/sensitivity.csv")
crit = read_csv("04_testing/results/criteria_assessment_full.csv")

rr = load_workbook(ROOT / "02_risk/risk_register.xlsx")["Risk Register"]
risks = []
for r in range(2, rr.max_row + 1):
    s = rr.cell(r, 4).value * rr.cell(r, 5).value
    risks.append(dict(id=rr.cell(r, 1).value, desc=rr.cell(r, 2).value, cat=rr.cell(r, 3).value, score=s, rating="High" if s >= 15 else "Medium" if s >= 8 else "Low",
                      nist=rr.cell(r, 8).value, nref=rr.cell(r, 9).value, iso=rr.cell(r, 10).value, uk=rr.cell(r, 12).value))
rc = Counter(r["rating"] for r in risks)
cm = load_workbook(ROOT / "03_controls/control_matrix.xlsx")["Control Matrix"]
controls = [dict(id=cm.cell(r, 1).value, risks=cm.cell(r, 2).value, desc=cm.cell(r, 3).value, inplace=cm.cell(r, 6).value, op=cm.cell(r, 11).value) for r in range(2, cm.max_row + 1)]
cc = Counter(c["inplace"] for c in controls)
cover = {}
for rk in risks:
    ctl = [c for c in controls if re.search(rk["id"] + r"\b", c["risks"] or "")]
    y = any(c["inplace"] == "Y" for c in ctl); p = any(c["inplace"] == "Partial" for c in ctl)
    cover[rk["id"]] = "no control identified" if not ctl else "covered" if y else "partly covered" if p else "gap (control not in place)"
covc = Counter(cover.values())
high_gap = [r["id"] for r in risks if r["rating"] == "High" and cover[r["id"]].startswith("gap")]

# ---------- page furniture ----------
def furniture(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5); canvas.setFillColor(GREY)
    canvas.drawString(18 * mm, 10 * mm, "Northbridge Assurance Review - self-assessment using audit methodology, not an independent audit. Subject: commit 9c4127a.")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
    canvas.setStrokeColor(colors.HexColor("#C9CED6")); canvas.line(18 * mm, 14 * mm, A4[0] - 18 * mm, 14 * mm)
    canvas.restoreState()


W = A4[0] - 36 * mm
S = []

# ---------- cover ----------
S += [Spacer(1, 30 * mm), P("Northbridge Assurance Review", title), P("Assurance report on the Northbridge Ledger Reconciliation &amp; Anomaly Detection Agent", sub), Spacer(1, 6 * mm)]
S.append(tbl([["Subject", "Northbridge Ledger Agent, commit 9c4127ad65aa5a9e7d362c2d05f69ecfdbf7c5f9 (23 September 2026)"],
              ["Report date", "23 September 2026"],
              ["Nature", "<b>Self-assessment using audit methodology.</b> The reviewer is also the developer of the system. This is not an independent audit and gives no assurance opinion or certification."],
              ["Challenge", "Testing was challenged by a separate Claude session (model-based, not a human reviewer). See section 2 and Appendix D."],
              ["Data", "Synthetic and vendor sandbox data only. No real client data."]],
             [32 * mm, W - 32 * mm], header=False, zebra=False))
S += [Spacer(1, 8 * mm), P("How to read this report", h2),
      P("Section 1 gives the conclusion and the numbers a decision maker needs. Sections 2 to 7 give scope, method, results, controls, findings and remediation. Appendices give the framework mapping, the control matrix extract, the evidence index and the challenge review. "
        "All figures come from a constructed test set designed by the reviewer after reading the rules; they describe how the system behaves on that set and are <b>not estimates of accuracy on real books</b>."),
      PageBreak()]

# ---------- 1 executive summary ----------
S += [P("1. Executive summary", h1)]
S.append(P("<b>Conclusion.</b> As at commit 9c4127a, Northbridge should <b>not be relied on as a control</b> for month-end reconciliation or anomaly review. It works as a demonstration and as a lead generator for the error types it was built to find, and it is fast and transparent. "
           "But it missed a whole error class (duplicate supplier bills), its reconciliation output changes with the order records are loaded, it keeps no record of what it ran, and it has no step where a person accepts or rejects a flag. "
           "On the eight criteria set before testing, seven were not met. A finance director could use it to prompt questions; they should not use a clean run as evidence that a ledger is free of duplicates or breaks.", callout))
S.append(P("What was reviewed", h2))
S.append(P("Northbridge reconciles bank lines against invoices and bills from Xero and QuickBooks and flags seven kinds of anomaly (A1 to A7). It is titled an \"Agent\" and positioned as a small-scale version of AI-native ledger products, but <b>calls no LLM</b>: it is deterministic rules and a median/MAD statistical check, with templated variance sentences. "
           "The review therefore treated it as a rules-based automated control and marked LLM-specific risks (prompt injection, hallucination, provider privacy) as not applicable."))
S.append(P("Key results", h2))
S.append(tbl([["Measure", "Result", "Criterion", "Assessment"],
              ["Recall, duplicates and reconciliation breaks", "83.3% (15/18)", ">= 95%", "Not met"],
              ["Recall, other rules (VAT, threshold, weekend, outlier)", "100% (13/13)", ">= 85%", "Met on point estimate; interval spans threshold"],
              ["Precision", "56.0% per item (53.7% per flag)", ">= 80%", "Not met"],
              ["F1", "0.691", ">= 0.85", "Not met"],
              ["False positive rate", "11.1% (22/199)", "<= 10%", "Not met (borderline, label-sensitive)"],
              ["Determinism (same data, any row order)", "Differs in 6 of 7 orderings tested", "100%", "Not met"],
              ["Audit trail completeness", "0% (no run log)", "100%", "Not met"],
              ["Human disposition of flags", "0% (no such step)", "100%", "Not met"]],
             [62 * mm, 44 * mm, 22 * mm, W - 128 * mm]))
S.append(P("Top findings", h2))
S.append(P("<b>F-01</b> Duplicate supplier bills are never detected (0 of 3): the duplicate check covers sales invoices only. <b>F-02</b> The reconciliation matcher takes the first match it finds, so results depend on record order and a genuine duplicate payment can vanish. "
           "<b>F-03</b> No run log. <b>F-04</b> No human sign-off, owner or intended-use statement. <b>F-05</b> The published 7/7 accuracy result is circular, and the public dashboard shows a hard-coded score. "
           "Eight further findings (6 Medium, 2 Low) are in section 6."))
S.append(P("Remediation", h2))
S.append(P("The Northbridge project was <b>not changed</b>. Fixes for eight findings (two only partly) were made in a scratch copy and delivered as a patch. On the same frozen dataset, duplicate supplier bills went from 0/3 to 3/3, output became identical across all seven row orderings, and variance sentences naming an opposing driver fell from 17 of 52 to 0 of 59. "
           "Two costs are disclosed: one duplicate payment (BT-044) is no longer flagged as such, and the project's published synthetic precision falls from 0.778 to 0.712. The fixes were designed against the same dataset used to retest, so they show the fixes work on those cases, not that real-world accuracy has improved."))
S.append(P("What a decision maker should take from this", h2))
S.append(P("Use it as a demonstration or an additional analytic. If it is to be used in a close process, first apply the patch, add a disposition step and run log, and test it on data the rules were not written to fit. "
           "An external auditor would take no controls reliance on it in its current form (section 7 and <i>fs_audit_implications.md</i>)."))
S.append(PageBreak())

# ---------- 2 scope ----------
S += [P("2. Scope, criteria, approach, independence and limitations", h1)]
S.append(P("Objective and scope", h2))
S.append(P("Assess whether Northbridge is governed, controlled and performs reliably enough to support month-end reconciliation, treating it as an automated control within an IT application. In scope: the Xero and QuickBooks connectors, data loading, reconciliation matcher, anomaly rules A1 to A7, variance explanations, the Streamlit dashboard, oversight, logging, change control, access, secrets and transparency. "
           "Out of scope: Sage or other connectors, production or real client data, LLM-specific tests (no LLM exists), penetration testing and Streamlit Cloud platform security."))
S.append(P("Reference frameworks", h2))
S.append(P("NIST AI RMF 1.0 as a governance lens; ISO/IEC 42001 Annex A, 23894 and 42005 where relevant; PRA SS1/23 for validation and change control; the five UK AI principles; OWASP LLM Top 10 considered and marked not applicable; ISA 315 (Revised 2019), 330, 500 and 265 for financial statement audit implications. "
           "Framework sub-references were cited at area level and must be checked against the source documents before reliance. EU AI Act dates from the project plan were not verified and are not relied on."))
S.append(P("Criteria (fixed before testing)", h2))
S.append(P("Eight criteria were set in the Terms of Reference before any test: recall (>= 95% for duplicates and breaks, >= 85% for other rules), precision >= 80%, F1 >= 0.85, false positive rate <= 10%, determinism 100%, audit trail completeness 100% and human disposition 100%. "
           "They are this review's criteria, not an industry standard; in a client engagement they would be agreed with the system owner. Recall is weighted above precision because a missed error usually costs more than a false alarm."))
S.append(P("Approach", h2))
S.append(P("Planning (terms of reference, system description); risk assessment (28 risks scored for likelihood and impact, mapped to frameworks); control matrix (29 controls, gaps recorded, not fixed before testing); testing (a 230-item synthetic ledger with ground-truth labels frozen by hash before the first run, an evaluation harness run against a read-only copy of the pinned commit, control tests, and a challenge review); findings in Condition, Criteria, Cause, Effect, Recommendation form; remediation demonstrated on a copy and retested."))
S.append(P("Independence disclosure", h2))
S.append(P("The reviewer developed the system, and the rules were built largely with an AI coding tool. Mitigations: criteria fixed before testing, gaps recorded as findings, hashes recorded before the first run (self-attested), and a challenge review by a separate Claude session given the raw evidence and not the reviewer's reasoning. "
           "That review is model-based and is not a substitute for an independent human review, which has not been done."))
S.append(P("Outcome of the challenge review", h2))
S.append(P("The challenge reviewer reproduced every headline metric exactly and raised 12 challenges (3 High, 8 Medium, 1 Low), <b>all accepted</b>. The most important: output depends on database row order (verified, and it withdrew an earlier \"determinism met\" result); the test set was built inside the rules' detection envelope, so the Terms of Reference wording \"not written to fit the rules\" was wrong and was amended; "
           "several labels (weekend, threshold, outlier) rest on narrative the data cannot show; and some control results were overstated and were reworded. Full disposition: Appendix D."))
S.append(P("Limitations", h2))
S.append(P("Synthetic and sandbox data only; one reviewer; point-in-time assessment of one commit; 3 to 6 errors per type, so per-type figures carry wide uncertainty; the 22 false positives arise from about six root causes, so simple confidence intervals are too narrow; "
           "connector ingestion against live systems, multi-line bills, voided or null data, multi-currency and several rule variations were not tested; the live Xero trial had lapsed, so pagination was not re-tested live; QuickBooks has no answer key, so no accuracy claim is made for it."))
S.append(PageBreak())

# ---------- 3 system overview ----------
S += [P("3. System overview and regulatory classification", h1)]
S.append(P("How Northbridge works", h2))
S.append(P("Connectors pull ledger data from Xero (read-only scopes) and QuickBooks (a read/write scope is granted, but the code only reads) into a local SQLite database. A matcher pairs each standalone bank line with an invoice or bill by amount, date and contact. "
           "Six rule checks and one statistical check then flag anomalies, and a template writes plain-English variance sentences. A Streamlit dashboard displays the results. The public deployment runs on frozen snapshots and makes no live API calls."))

# data flow drawing
d = Drawing(W, 62 * mm)
def box(x, y, w, h, txt, fill="#E8EEF7", stroke="#1F3864", size=7.5):
    d.add(Rect(x, y, w, h, fillColor=colors.HexColor(fill), strokeColor=colors.HexColor(stroke), strokeWidth=0.8))
    for i, ln in enumerate(txt.split("\n")):
        d.add(String(x + w / 2, y + h - 10 - i * 9, ln, textAnchor="middle", fontName="Helvetica", fontSize=size))
def arrow(x1, y1, x2, y2):
    d.add(Line(x1, y1, x2, y2, strokeColor=colors.HexColor("#595959"), strokeWidth=1))
    d.add(Line(x2, y2, x2 - 4 * (1 if x2 > x1 else -1 if x2 < x1 else 0), y2 + (3 if x2 == x1 else 0), strokeColor=colors.HexColor("#595959"), strokeWidth=1))
d.add(String(0, 62 * mm - 8, "Data flow (red = data leaves the developer machine)", fontName="Helvetica-Bold", fontSize=8))
box(0, 100, 70, 50, "Xero API\n(read scopes)", "#FDECEA", "#B71C1C")
box(0, 30, 70, 50, "QuickBooks API\n(read/write scope\ngranted)", "#FDECEA", "#B71C1C")
box(105, 55, 80, 70, "Connectors\nOAuth + GET\n.env holds keys\nand plaintext\nrefresh tokens")
box(220, 55, 70, 70, "SQLite ledger\n(local)\ndemo copies\ncommitted")
box(320, 55, 90, 70, "Matcher, rules\nA1-A7,\nvariance\ntemplates")
box(440, 55, 60, 70, "Dashboard\n(display only,\nno sign-off)")
box(320, 0, 90, 34, "Public GitHub +\nStreamlit Cloud", "#FDECEA", "#B71C1C")
for a in [(70, 125, 105, 100), (70, 55, 105, 80), (185, 90, 220, 90), (290, 90, 320, 90), (410, 90, 440, 90)]:
    arrow(*a)
arrow(255, 55, 340, 34)
S.append(d)
S.append(P("No LLM provider receives data. Personal data risk is low today because the data is fictional or vendor sandbox data (the Xero demo uses real brand names as counterparties; the sample of contact names checked was limited)."))
S.append(P("Regulatory classification", h2))
S.append(tbl([["Question", "Conclusion"],
              ["AI system under the EU AI Act?", "Probably not in the default configuration (fixed rules and classical statistics). Borderline if the optional IsolationForest path is enabled."],
              ["Prohibited practice; Annex III high-risk use; Article 50 transparency", "No; not a listed use; not triggered"],
              ["Risk tier", "Outside scope / minimal risk on the current design"],
              ["Reassess if", "an LLM or trained model is added, real client data is used, decisions about individuals are made, or the tool is placed on the EU market"]],
             [58 * mm, W - 58 * mm]))
S.append(Spacer(1, 3 * mm))
S.append(P("Against the five UK principles: <b>safety, security and robustness</b> partly met; <b>transparency and explainability</b> partly met (rules are inspectable, but the \"Agent\" title and \"AI-native\" positioning can be read as AI, and the README never says no model is used); <b>fairness</b> largely not relevant; "
           "<b>accountability and governance</b> not met (no owner, sign-off or audit trail); <b>contestability and redress</b> not met (no way to record that a flag was wrong or trace a result to a version). "
           "This is the reviewer's reasoning, not legal advice; legal status must be re-checked before it is cited."))
S.append(PageBreak())

# ---------- 4 results ----------
S += [P("4. Results against criteria", h1)]
S.append(P("Test set: 230 synthetic items (bills, sales invoices, bank lines) in Northbridge's own schema, with 31 real errors across eight types and 199 clean items, of which 32 are \"hard\" legitimate items that look like errors (bank charges, weekend suppliers, fixed-price bills, a supplier that became VAT-registered, large one-off purchases). "
           "Labels were frozen by hash before the first run. The set is a constructed mix; precision and false positive rates depend on how many hard cases were included (1.2% of standard clean items were flagged, 62.5% of hard ones)."))
rows = [["Criterion", "Threshold", "Result", "Assessment"]]
for c in crit:
    rows.append([c["criterion"], c["threshold"], c["result"], c["assessment"]])
S.append(tbl(rows, [50 * mm, 17 * mm, 52 * mm, W - 119 * mm]))
S.append(P("Item level, any rule: 28 true positives, 3 false negatives, 22 false positives, 177 true negatives. Recall 0.903 (95% interval 0.751 to 0.967), precision 0.560 (0.423 to 0.688), false positive rate 0.111 (0.074 to 0.162). Intervals assume independent items and are too narrow because false positives cluster in about six causes.", small))

# recall chart
types = ["DUP_AR", "DUP_AP", "DUP_PAY", "BREAK", "VAT_MISCODE", "THRESHOLD", "WEEKEND", "OUTLIER"]
names = {"DUP_AR": "Duplicate sales invoice", "DUP_AP": "Duplicate supplier bill", "DUP_PAY": "Duplicate payment", "BREAK": "Reconciliation break", "VAT_MISCODE": "VAT miscoding", "THRESHOLD": "Just under threshold", "WEEKEND": "Weekend-dated bill", "OUTLIER": "Outlier amount"}
ch = Drawing(W, 46 * mm)
ch.add(String(0, 46 * mm - 9, "Recall by error type: before (grey) and after remediation (blue); dashed line = criterion", fontName="Helvetica-Bold", fontSize=8))
x0, wmax, y = 92 * mm, W - 92 * mm - 22 * mm, 46 * mm - 22
for t in types:
    b, a = float(before_t[t]["recall_any_rule"]), float(after_t[t]["recall_any_rule"])
    nn = before_t[t]["n_errors"]
    ch.add(String(x0 - 4, y + 4, f"{names[t]} (n={nn})", textAnchor="end", fontName="Helvetica", fontSize=7.5))
    ch.add(Rect(x0, y + 7, max(b * wmax, 0.5), 5, fillColor=colors.HexColor("#9AA3AF"), strokeColor=None))
    ch.add(Rect(x0, y, max(a * wmax, 0.5), 5, fillColor=colors.HexColor("#1F3864"), strokeColor=None))
    ch.add(String(x0 + max(a, b) * wmax + 4, y + 3, f"{b:.0%} / {a:.0%}", fontName="Helvetica", fontSize=7))
    thr = 0.95 if t in ("DUP_AR", "DUP_AP", "DUP_PAY", "BREAK") else 0.85
    ch.add(Line(x0 + thr * wmax, y - 2, x0 + thr * wmax, y + 14, strokeColor=colors.black, strokeDashArray=[2, 2], strokeWidth=0.7))
    y -= 13.5
S.append(ch)
S.append(P("Sensitivity to labelling (challenge review)", h2))
S.append(P("Precision and false positive rate depend on labelling choices, so they are shown under alternatives. Recall on duplicates and breaks fails only because of the supplier-bill gap, a code limit, not a labelling effect.", body))
srows = [["Scenario", "FP", "Precision", "F1", "FPR", "Recall"]]
for r in sens:
    srows.append([r["scenario"], r["FP"], r["precision"], r["F1"], r["FPR"], r["recall"]])
S.append(tbl(srows, [W - 92 * mm, 12 * mm, 20 * mm, 18 * mm, 20 * mm, 22 * mm]))
S.append(P("<b>Determinism:</b> five fresh-process runs on one database were identical, but a different row order changed the result in 6 of 7 orderings (one lost a true duplicate payment), so the criterion is not met. <b>Published results:</b> the project's 7/7 and 84/84 figures reproduce exactly but use planted anomalies written to fit the rules. Engine speed: about 0.01 s on 230 items.", small))
S.append(PageBreak())

# ---------- 5 control assessment ----------
S += [P("5. Control assessment summary", h1)]
S.append(P(f"The 28 risks were scored before controls: <b>{rc['High']} High, {rc['Medium']} Medium, {rc['Low']} Low</b> (list in Appendix A). "
           f"{len(controls)} controls were identified from the code and documents: <b>{cc['Y']} in place, {cc['Partial']} partial, {cc['N']} not in place</b>. Controls that are not in place were recorded as gaps and were not added before testing. "
           f"Coverage of the 28 risks: {covc['covered']} covered, {covc['partly covered']} partly covered, {covc['gap (control not in place)']} with a control identified but not in place, {covc['no control identified']} with no control identified (R22, vendor and trial changes, accepted as Low). "
           f"High risks with no working control: {', '.join(high_gap)} (no human disposition step; no run log)."))
S.append(P("Controls that operated as designed", h2))
S.append(P("Xero requests read scopes only and neither connector makes write calls; the local secrets file is git-ignored and was never committed; no credential patterns were found in code files across the full history (a pattern search, not a scanner run); the dashboard's data-quality counts equal an independent recount; "
           "every flag carries a reason string; the published accuracy figures reproduce; fixed random seeds work for a given input order."))
S.append(P("Controls that failed or are absent (operating result)", h2))
pick = ["C03", "C05", "C10", "C13", "C15", "C20", "C21", "C22", "C24", "C25", "C26", "C27", "C28"]
rows = [["ID", "Control", "In place", "Operating result"]]
for c in controls:
    if c["id"] in pick:
        rows.append([c["id"], c["desc"], c["inplace"], c["op"]])
S.append(tbl(rows, [10 * mm, 60 * mm, 14 * mm, W - 84 * mm]))
S.append(P("Full matrix with design tests, planned operating tests, evidence references and results: <i>03_controls/control_matrix.xlsx</i>. Controls C12 (pagination against live systems) and C29 (independent challenge) could not be fully tested; see Appendix C.", small))
S.append(PageBreak())

# ---------- 6 findings ----------
S += [P("6. Findings and recommendations", h1)]
S.append(P("Ratings: <b>High</b> can let a material error go undetected or unevidenced, or makes results unreliable; <b>Medium</b> weakens reliability, control or transparency and matters beyond a demo; <b>Low</b> is minor. Findings describe the system at the pinned commit. Full 5-C write-ups: <i>05_reporting/findings.md</i>."))
F = [
 ("F-01", "High", "Duplicate supplier bills are not detected",
  "0 of 3 duplicated supplier bills flagged; duplicate sales invoices 5 of 5.",
  "Recall on duplicates and breaks >= 95%; result 83.3%.",
  "Rule A1 selects sales invoices only (rules.py).",
  "Bills entered twice go unreported; duplicate payables are a principal cash-loss risk.",
  "Run A1 over both sales invoices and supplier bills; add regression cases."),
 ("F-02", "High", "Reconciliation output depends on record order and mis-attributes recurring payments",
  "Same data in another row order gave different flags (6 of 7 orderings differ), including loss of a true duplicate payment; two legitimate receipts were attributed to an already-paid invoice of the same amount and flagged as duplicate payments.",
  "Identical output for identical data; documents matched correctly.",
  "Matcher takes the first candidate within tolerance: no ordering, no closest-date preference, no one-to-one allocation (matcher.py).",
  "Errors can appear or vanish with no change in the ledger. Shown on constructed data; mechanism is in the code.",
  "Deterministic candidate ranking, exclude matched or paid invoices, one-to-one allocation, row-order test."),
 ("F-03", "High", "No run log or audit trail",
  "No logging, input hash, code version or timestamp; schema has no audit table.",
  "100% of runs traceable to input and code version.",
  "Not designed in.",
  "Results cannot be evidenced or re-performed; prerequisite for any reliance is missing.",
  "Append-only run record: input hash, git SHA, dependency versions, timestamp, user, output hash."),
 ("F-04", "High", "No human disposition of flags; no owner or intended-use statement",
  "Dashboard displays flags with no accept, reject or reason step; owner and intended use undocumented.",
  "100% of flags dispositioned before close; documented accountability.",
  "Display-only design.",
  "Oversight cannot be evidenced; no way to record that a flag was wrong or handled.",
  "Disposition record per flag tied to the run log; document owner and intended use."),
 ("F-05", "High", "Reported accuracy rests on circular validation and a static score",
  "7/7 and 84/84 reproduce but use anomalies written to fit the rules; public dashboard shows a hard-coded score when the answer key is absent.",
  "Accuracy claims evidenced on data not written to fit the rules, with clean items.",
  "Validation designed to show mechanics; fallback score is a constant.",
  "Readers may take 7/7 as evidence of general accuracy.",
  "Restate as mechanics tests; label the fallback as static; build an independent test set before any accuracy claim."),
]
for f in F:
    rows = [[P(f"<b>{f[0]}  {f[2]}</b>", cellb), P(f"<b>{f[1]}</b>", cellb)],
            [P("<b>Condition</b>", cell), P(f[3], cell)], [P("<b>Criteria</b>", cell), P(f[4], cell)], [P("<b>Cause</b>", cell), P(f[5], cell)],
            [P("<b>Effect</b>", cell), P(f[6], cell)], [P("<b>Recommendation</b>", cell), P(f[7], cell)]]
    t = Table(rows, colWidths=[24 * mm, W - 24 * mm])
    t.setStyle(TableStyle([("SPAN", (0, 0), (0, 0)), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FDECEA")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED6")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
    # header row spans both columns
    t = Table([[P(f"<b>{f[0]}  {f[2]}</b>  -  <font color='#B71C1C'><b>{f[1]}</b></font>", cell), ""]] + rows[1:], colWidths=[24 * mm, W - 24 * mm])
    t.setStyle(TableStyle([("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FDECEA")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED6")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
    S += [KeepTogether([t, Spacer(1, 4)])]
S.append(P("Medium and Low findings", h2))
M = [["ID", "Rating", "Finding and condition", "Recommendation"],
     ["F-06", "Medium", "<b>No regression tests, version tags or pinned dependencies.</b> No tests directory, 13 commits, 0 tags, unpinned requirements. Any edit can change detections unnoticed.", "Regression suite, tagged releases, pinned dependencies, version in output."],
     ["F-07", "Medium", "<b>Variance explanations can point against the headline.</b> 17 of 52 test sentences (35 of 69 on the project's synthetic data) name a driver moving opposite to the headline; contacts that disappeared are never named. Figures agree with source; text is templated, not hallucinated.", "Rank drivers in the headline's direction; include disappeared contacts; state share explained."],
     ["F-08", "Medium", "<b>False positives on legitimate items exceed the criterion.</b> 22 of 199 clean items (bank charges 6, fixed-price and boundary bills 5, weekend suppliers 4, VAT registration change 3, large equipment 2, recurring receipts 2). Label-sensitive.", "Allow-lists for recurring fixed charges and known weekend suppliers; bank-charge handling; record accepted exceptions."],
     ["F-09", "Medium", "<b>Excess QuickBooks permission; plaintext refresh tokens.</b> Read/write scope granted though code only reads; tokens written in plaintext to a local file.", "Read-only scope if offered, else test for no write calls; keychain or secrets manager; rotate tokens."],
     ["F-10", "Medium", "<b>Ambiguous \"AI-native agent\" positioning; no limitations statement.</b> The README positions the tool as a small-scale version of AI-native products and the title says \"Agent\"; no LLM is used and the README never says so. It omits intended use, limits on real books and production status.", "Describe as rules-based; add intended use and limitations."],
     ["F-11", "Medium", "<b>Data completeness not evidenced.</b> Pagination tested with mocks only; no validation gate; coverage table not reconciled to source.", "Reconcile loaded counts and totals to the source system; test ingestion live."],
     ["F-12", "Low", "<b>Hard-coded, unapproved thresholds and message defects.</b> Constants have no rationale or approver; one message reads \"Bank line for None\".", "Documented configuration with approver; fix messages."],
     ["F-13", "Low", "<b>QuickBooks bank purchases all reported as unmatched.</b> 0 of 40 auto-matched; the sandbox has no bills for direct expenses (noise, not a shown matcher failure).", "Treat direct expenses as a separate class or state the limitation."]]
S.append(tbl(M, [11 * mm, 15 * mm, W - 92 * mm, 66 * mm]))
S.append(PageBreak())

# ---------- 7 remediation ----------
S += [P("7. Remediation results and financial statement audit implications", h1)]
S.append(P("Remediation (demonstrated on a copy)", h2))
S.append(P("The Northbridge project was not changed. The pinned commit was copied, fixes were applied to the copy and captured in <i>05_reporting/remediation.patch</i>, which applies cleanly to a fresh checkout of the pinned commit; its 10 new unit tests pass there. The owner decides whether to apply it. "
           "Changes: A1 covers supplier bills; deterministic one-to-one matcher; optional run log (off by default); pinned dependencies, regression tests and a version; variance drivers follow the headline's direction; README rewritten as rules-based with limitations; parameters moved to a config file and messages fixed. "
           "Not changed: F-04, F-08, F-09, F-11 and F-13, which need design or process decisions, vendor options or live credentials."))
S.append(tbl([["Measure (same frozen dataset and harness)", "Before", "After", "Criterion"],
              ["Duplicate supplier bills detected", "0/3", "3/3", ""],
              ["Recall, duplicates and breaks", "0.833 (15/18)", "0.944 (17/18)", ">= 0.95, still not met by one item"],
              ["Precision (per item)", "0.560", "0.600", ">= 0.80, not met"],
              ["F1", "0.691", "0.741", ">= 0.85, not met"],
              ["False positive rate", "0.111 (22/199)", "0.101 (20/199)", "<= 0.10, not met (10.05%)"],
              ["False positives on standard clean items", "2/167", "0/167", ""],
              ["Identical output across 7 row orderings", "No (6 of 7 differ)", "Yes (7 of 7)", "100%"],
              ["Variance sentences with an opposing driver", "17 of 52", "0 of 59", ""],
              ["Published synthetic score", "84/84, 24 FP, precision 0.778", "84/84, 34 FP, precision 0.712", "see below"]],
             [58 * mm, 36 * mm, 36 * mm, W - 130 * mm]))
S.append(Spacer(1, 2 * mm))
S.append(P("<b>Costs and limits, stated plainly.</b> (1) The true duplicate payment BT-044 is no longer flagged as a duplicate payment: its bank line now attaches to an unpaid bill of the same supplier and amount, which in this dataset is the deliberately planted duplicate of the paid bill; the tool flags that bill as a duplicate instead. The label was not changed after seeing the result. "
           "(2) Ten new supplier-bill duplicate flags on the project's synthetic data count as false positives against its planted list, lowering its published precision. (3) The fixes were designed while looking at this dataset, so the retest shows they work on those cases and does not show accuracy on real books has improved. After remediation two of eight criteria are met (determinism; recall on other rules on point estimate)."))
S.append(P("Implications for a financial statement audit", h2))
S.append(P("Treated as part of an entity's information system (ISA 315), an auditor would take <b>no controls reliance</b>:"))
S.append(tbl([["Layer", "Result", "Effect"],
              ["IT general controls (ISA 315)", "No run log, no tests, tags or pinning, plaintext tokens and excess scope", "Not effective; prevents reliance on automated output"],
              ["Tool as an automated control (ISA 330)", "Duplicates and breaks recall 83.3% vs 95%; output varies with record order", "Does not meet the precision needed for payables completeness and accuracy"],
              ["Human review control", "No disposition step, owner or intended use", "No management review control to test"],
              ["Substantive procedures (ISA 330, 500)", "Exceptions list is order-dependent and misses supplier duplicates", "Auditor re-performs from source data, for example a duplicate analysis over the full payables population"]],
             [40 * mm, 62 * mm, W - 102 * mm]))
S.append(P("Candidates for written communication of significant deficiencies to those charged with governance (ISA 265), if an entity relied on the tool: F-03 (no audit trail), F-04 (no evidenced review) and F-01 (undetected duplicate payables). Significance depends on the entity and the auditor's judgement. Even with better recall, absent logging and change control would prevent reliance.", body))
S.append(PageBreak())

# ---------- Appendices ----------
S += [P("Appendix A. Risk register and framework mapping", h1)]
S.append(P("28 risks, each scored 1 to 5 for likelihood and impact before controls (inherent score = product; High >= 15, Medium 8 to 14, Low < 8). NIST and ISO references are at area level and must be verified against the source texts. Full register with rationale: <i>02_risk/risk_register.xlsx</i>.", small))
cats = []
for r in risks:
    if r["cat"] not in cats:
        cats.append(r["cat"])
rows = [["Category", "Risks", "High", "Medium", "Low"]]
for c in cats:
    rs = [r for r in risks if r["cat"] == c]
    rows.append([c, str(len(rs)), str(sum(r["rating"] == "High" for r in rs)), str(sum(r["rating"] == "Medium" for r in rs)), str(sum(r["rating"] == "Low" for r in rs))])
rows.append(["<b>Total</b>", f"<b>{len(risks)}</b>", f"<b>{rc['High']}</b>", f"<b>{rc['Medium']}</b>", f"<b>{rc['Low']}</b>"])
S.append(tbl(rows, [70 * mm, 20 * mm, 20 * mm, 20 * mm, 20 * mm]))
S.append(Spacer(1, 3 * mm))
S.append(P("High-rated risks", h2))
rows = [["ID", "Risk", "Score", "NIST function / ref", "ISO/IEC 42001 area", "UK principle"]]
for r in [x for x in risks if x["rating"] == "High"]:
    rows.append([r["id"], r["desc"], str(r["score"]), f"{r['nist']} / {r['nref']}", r["iso"], r["uk"]])
S.append(tbl(rows, [10 * mm, 62 * mm, 11 * mm, 30 * mm, 33 * mm, W - 146 * mm]))
S.append(P("Risks recorded as not applicable because no LLM is called: prompt injection, hallucinated explanations, provider data handling and retention, silent model changes, model and prompt inventory.", small))
S.append(PageBreak())

S += [P("Appendix B. Control matrix extract", h1)]
S.append(P("Design and operating assessment for a selection of controls. In place: Y = design adequate; Partial = design limited; N = not in place.", small))
pickB = ["C01", "C02", "C04", "C06", "C11", "C14", "C16", "C17", "C19", "C22", "C23"]
rows = [["ID", "Control", "In place", "Operating result"]]
for c in controls:
    if c["id"] in pickB:
        rows.append([c["id"], c["desc"], c["inplace"], c["op"]])
S.append(tbl(rows, [10 * mm, 58 * mm, 14 * mm, W - 82 * mm]))
S.append(PageBreak())

S += [P("Appendix C. Evidence index", h1)]
S.append(P("All files are in the review repository. The subject is pinned to commit 9c4127ad65aa5a9e7d362c2d05f69ecfdbf7c5f9. Test dataset and labels were hashed before the first run (self-attested).", small))
E = [["Ref", "What", "Where"],
     ["E-01", "Test dataset and labels frozen (SHA-256 recorded)", "04_testing/evidence_log.md; test_dataset/"],
     ["E-02", "Evaluation harness run: metrics, per-type and per-rule results", "04_testing/results/"],
     ["E-03", "Determinism (5 fresh processes); withdrawn and corrected in E-06", "results/run_metadata.json; results/order_sensitivity.txt"],
     ["E-04", "Control operating tests (scopes, secrets, coverage counts, published-figure reproduction, documentation checks)", "results/control_tests_output.txt"],
     ["E-05", "Variance direction check", "results/variance_direction_check.txt"],
     ["E-06", "Challenge review outcome and corrections; order sensitivity, sensitivity analysis, full criteria", "04_testing/challenge_review.md; challenge_review_raw.md; results/sensitivity.csv; results/criteria_assessment_full.csv"],
     ["E-07", "Remediation demonstrated on a copy; retest", "05_reporting/remediation.md, remediation.patch; results_remediated/"],
     ["-", "Terms of reference (with Amendment 1), system description", "01_planning/"],
     ["-", "Risk register; regulatory classification", "02_risk/"],
     ["-", "Control matrix", "03_controls/control_matrix.xlsx"],
     ["-", "Findings; financial statement audit implications", "05_reporting/findings.md; fs_audit_implications.md"]]
S.append(tbl(E, [14 * mm, W - 110 * mm, 96 * mm]))
S.append(Spacer(1, 3 * mm))
S.append(P("Not tested: connector ingestion against live systems and pagination (credentials and the Xero trial unavailable); scopes actually granted in the vendor portals; multi-line bills, voided or null data, multi-currency; further rule variations; real-book accuracy. A human independent review has not been done.", small))

S.append(P("Appendix D. Challenge review: summary of disposition", h1))
D = [["ID", "Sev", "Challenge", "Outcome"],
     ["CH-01", "High", "Determinism overstated: output depends on row order", "Accepted, re-verified; criterion re-rated not met; finding F-02"],
     ["CH-02", "High", "Test set built inside the rules' envelope", "Accepted; Terms of Reference amended; results reported as a constructed mix"],
     ["CH-03", "High", "Weekend, threshold, outlier labels rest on unobservable narrative", "Accepted; alternative labelling shown in sensitivity analysis"],
     ["CH-04", "Med", "Criterion outcomes flip under relabelling", "Accepted; false positive rate reported as borderline"],
     ["CH-05", "Med", "Confidence intervals too narrow (clustered false positives)", "Accepted; caveat added"],
     ["CH-06", "Med", "Only 5 of 8 criteria scored; precision per item not per flag", "Accepted; all 8 scored; per-flag 0.537 reported"],
     ["CH-07 to 09", "Med", "Control results overstated (C22, C19, C06, C07, C11, C23)", "Accepted; wording corrected in the matrix"],
     ["CH-10", "Med", "\"Likely on real data\" is extrapolation", "Accepted; reworded as demonstrated on constructed data"],
     ["CH-11", "Med", "Coverage not tested (ingestion, multi-line bills, null data)", "Accepted as a scope limit"],
     ["CH-12", "Low", "Frozen-before-run is self-attested", "Accepted; labelled as such"]]
S.append(tbl(D, [20 * mm, 12 * mm, 68 * mm, W - 100 * mm]))

doc = SimpleDocTemplate(str(OUT), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=20 * mm,
                        title="Northbridge Assurance Review", author="Self-assessment", subject="Assurance report (self-assessment, not an independent audit)")
doc.build(S, onFirstPage=furniture, onLaterPages=furniture)
print("built", OUT)
